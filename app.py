"""ATS Resume Checker - Streamlit + Gemini Flash.

Upload a resume (PDF / DOCX / TXT), optionally paste a job description,
and get an ATS score with concrete improvement suggestions.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
DEFAULT_MODEL = "gemini-2.5-flash"
FALLBACK_MODEL = "gemini-flash-latest"  # alias that always points to newest Flash
MAX_FILE_MB = 5
MAX_RESUME_CHARS = 20_000
MAX_JD_CHARS = 8_000
MIN_RESUME_CHARS = 100

# Weights used to compute the overall score (must sum to 1.0)
WEIGHTS = {
    "keywords": 0.30,
    "content_impact": 0.25,
    "formatting": 0.20,
    "structure": 0.15,
    "readability": 0.10,
}
LABELS = {
    "keywords": "Keywords & Relevance",
    "content_impact": "Content & Impact",
    "formatting": "ATS Formatting",
    "structure": "Structure & Sections",
    "readability": "Readability & Grammar",
}

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume strictly and honestly. Do not \
inflate scores. Base every statement ONLY on the resume text provided; never \
invent experience, numbers, or skills the candidate does not have.

Return ONLY a valid JSON object (no markdown, no commentary) with this exact schema:
{
  "scores": {
    "keywords": <int 0-100>,
    "content_impact": <int 0-100>,
    "formatting": <int 0-100>,
    "structure": <int 0-100>,
    "readability": <int 0-100>
  },
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", ...],
  "missing_keywords": ["<string>", ...],
  "improvements": [
    {
      "priority": "High" | "Medium" | "Low",
      "section": "<resume section or 'General'>",
      "issue": "<what is wrong>",
      "fix": "<specific, actionable fix>",
      "example": "<optional improved rewrite of an actual line, or empty string>"
    }
  ]
}

Scoring guide:
- keywords: relevant hard/soft skills, tools, industry terms (match the job description if given).
- content_impact: quantified achievements, action verbs, results over duties.
- formatting: parsable by ATS (no tables/columns/graphics artifacts, standard fonts, consistent dates, contact info present).
- structure: standard sections (Contact, Summary, Experience, Education, Skills), logical order, appropriate length.
- readability: grammar, clarity, concise bullets, consistent tense.
Give 3-6 strengths, up to 10 missing keywords, and 5-10 improvements sorted by priority."""


# --------------------------------------------------------------------------- #
# File parsing
# --------------------------------------------------------------------------- #
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX or TXT bytes."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # Resumes often keep content in tables; include it
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload PDF, DOCX or TXT.")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# --------------------------------------------------------------------------- #
# Gemini helpers
# --------------------------------------------------------------------------- #
def parse_json_response(raw: str) -> dict:
    """Parse JSON from model output, tolerating markdown fences / stray text."""
    if not raw:
        raise ValueError("Empty response from the model.")
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise ValueError("Model did not return valid JSON.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Validate the model output, clamp scores and compute the overall score."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected response format.")
    raw_scores = data.get("scores") or {}
    scores = {k: _clamp(raw_scores.get(k)) for k in WEIGHTS}
    overall = _clamp(sum(scores[k] * w for k, w in WEIGHTS.items()))

    def str_list(key, limit):
        items = data.get(key) or []
        return [str(i).strip() for i in items if str(i).strip()][:limit]

    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section") or "General"),
                "issue": str(item.get("issue") or "").strip(),
                "fix": str(item.get("fix") or "").strip(),
                "example": str(item.get("example") or "").strip(),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda x: order[x["priority"]])

    return {
        "overall": overall,
        "scores": scores,
        "summary": str(data.get("summary") or "").strip(),
        "strengths": str_list("strengths", 8),
        "missing_keywords": str_list("missing_keywords", 15),
        "improvements": improvements[:12],
    }


def analyze_resume(api_key: str, model: str, resume_text: str, job_desc: str = "") -> dict:
    """Send the resume to Gemini and return a normalized analysis dict."""
    client = genai.Client(api_key=api_key)

    prompt = f"RESUME:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\"\n"
    if job_desc.strip():
        prompt += (
            f"\nTARGET JOB DESCRIPTION:\n\"\"\"\n{job_desc.strip()[:MAX_JD_CHARS]}\n\"\"\"\n"
            "\nScore keywords based on match with this job description."
        )
    else:
        prompt += "\nNo job description given; evaluate for general ATS-friendliness."

    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        temperature=0.2,
        response_mime_type="application/json",
    )

    models_to_try = [model] + ([FALLBACK_MODEL] if model != FALLBACK_MODEL else [])
    last_error = None
    for m in models_to_try:
        try:
            response = client.models.generate_content(model=m, contents=prompt, config=config)
            return normalize_result(parse_json_response(response.text))
        except Exception as e:  # try next model, then surface the error
            last_error = e
            msg = str(e).lower()
            # Bad key / quota problems won't be fixed by switching models
            if any(s in msg for s in ("api key", "api_key", "permission", "quota", "429")):
                break
    raise RuntimeError(f"Gemini request failed: {last_error}")


# --------------------------------------------------------------------------- #
# UI
# --------------------------------------------------------------------------- #
def get_secret(name: str) -> str:
    """Read from env var first, then Streamlit secrets (which may not exist locally)."""
    value = os.environ.get(name, "")
    if value:
        return value
    try:
        return st.secrets.get(name, "")
    except Exception:
        return ""


def score_color(score: int) -> str:
    return "🟢" if score >= 75 else "🟡" if score >= 50 else "🔴"


def render_results(result: dict) -> None:
    overall = result["overall"]
    verdict = (
        "Excellent - highly ATS-friendly" if overall >= 80
        else "Good - a few tweaks needed" if overall >= 65
        else "Fair - needs notable improvements" if overall >= 45
        else "Poor - major improvements needed"
    )
    st.subheader("Your ATS Score")
    c1, c2 = st.columns([1, 2])
    c1.metric("Overall", f"{overall} / 100")
    c2.markdown(f"**{score_color(overall)} {verdict}**")
    c2.progress(overall / 100)
    if result["summary"]:
        st.info(result["summary"])

    st.subheader("Score Breakdown")
    cols = st.columns(len(result["scores"]))
    for col, (key, val) in zip(cols, result["scores"].items()):
        col.metric(LABELS[key], f"{val}")
        col.progress(val / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["No strengths identified."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("🔑 Missing Keywords")
        if result["missing_keywords"]:
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.markdown("No major keyword gaps found.")

    st.subheader("🛠️ Recommended Improvements")
    icons = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}
    for imp in result["improvements"]:
        title = f"{icons[imp['priority']]} {imp['priority']} · {imp['section']}"
        with st.expander(title, expanded=imp["priority"] == "High"):
            st.markdown(f"**Issue:** {imp['issue']}")
            st.markdown(f"**Fix:** {imp['fix']}")
            if imp["example"]:
                st.markdown("**Example rewrite:**")
                st.code(imp["example"], language=None)

    st.download_button(
        "⬇️ Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tailored improvement tips, powered by Gemini Flash.")

    with st.sidebar:
        st.header("⚙️ Settings")
        api_key = get_secret("GEMINI_API_KEY")
        if api_key:
            st.success("API key loaded from secrets.")
        else:
            api_key = st.text_input(
                "Gemini API key", type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        model = st.text_input("Model", value=get_secret("GEMINI_MODEL") or DEFAULT_MODEL)
        st.caption("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    col_a, col_b = st.columns(2)
    with col_a:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col_b:
        job_desc = st.text_area(
            "Job description (optional, improves keyword matching)", height=170,
            placeholder="Paste the job description here...",
        )

    if st.button("🚀 Analyze Resume", type="primary"):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large (max {MAX_FILE_MB} MB).")
            return

        try:
            text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            return
        if len(text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text was found. If this is a scanned/image PDF, "
                "export a text-based PDF or upload a DOCX instead."
            )
            return

        with st.spinner("Analyzing your resume..."):
            try:
                st.session_state["result"] = analyze_resume(
                    api_key, model.strip() or DEFAULT_MODEL, text, job_desc
                )
            except Exception as e:
                st.session_state.pop("result", None)
                st.error(str(e))
                return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
